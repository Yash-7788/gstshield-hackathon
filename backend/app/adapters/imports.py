"""Bounded structured parsers. Invoked inside the disposable parser process."""

import csv
import io
import json
import re
from contextlib import suppress
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from zipfile import ZipFile

from defusedxml import ElementTree as SafeXML
from openpyxl import load_workbook

from app.domain.imports import FIELDS, REQUIRED, ParseFailure, canonical_row, flag_duplicates


def strict_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ParseFailure("DUPLICATE_JSON_KEY")
        result[key] = value
    return result


def bounded_json(text: str, depth: int, *, exact_numbers=False):
    level = 0
    quoted = escaped = False
    for char in text:
        if quoted:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                quoted = False
        elif char == '"':
            quoted = True
        elif char in "[{":
            level += 1
            if level > depth:
                raise ParseFailure("JSON_DEPTH_LIMIT")
        elif char in "]}":
            level -= 1
    try:
        return json.loads(
            text,
            object_pairs_hook=strict_object,
            parse_float=Decimal if exact_numbers else float,
            parse_constant=lambda _: (_ for _ in ()).throw(ParseFailure("INVALID_JSON")),
        )
    except (ValueError, RecursionError):
        raise ParseFailure("INVALID_JSON") from None


def check_columns(columns, limits):
    if not columns or len(columns) > limits["max_import_columns"]:
        raise ParseFailure("COLUMN_LIMIT")
    if any(
        not isinstance(c, str)
        or not c
        or len(c) > 128
        or any(ord(char) < 32 or ord(char) == 127 for char in c)
        for c in columns
    ):
        raise ParseFailure("INVALID_HEADERS")
    if len(set(columns)) != len(columns):
        raise ParseFailure("DUPLICATE_HEADERS")


def check_cell(value, limits):
    if value is not None and not isinstance(value, (str, int, float, date, datetime)):
        raise ParseFailure("INVALID_CELL")
    if isinstance(value, bool) or (
        isinstance(value, str) and len(value) > limits["max_cell_characters"]
    ):
        raise ParseFailure("CELL_LIMIT")


def csv_rows(content, limits):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeError:
        raise ParseFailure("UTF8_REQUIRED") from None
    if "\x00" in text:
        raise ParseFailure("INVALID_CSV")
    previous = csv.field_size_limit(limits["max_cell_characters"])
    try:
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        columns = next(reader, [])
        check_columns(columns, limits)
        rows = []
        for row in reader:
            if not row:
                row = ["" for _ in columns]
            if len(rows) >= limits["max_import_rows"]:
                raise ParseFailure("ROW_LIMIT")
            if len(row) != len(columns):
                raise ParseFailure("CSV_ROW_WIDTH_INVALID")
            for value in row:
                check_cell(value, limits)
            rows.append(dict(zip(columns, row, strict=True)))
        return columns, rows, None
    except csv.Error:
        raise ParseFailure("INVALID_CSV") from None
    finally:
        csv.field_size_limit(previous)


def inspect_xlsx(content, limits):
    """Read actual expanded bytes; declared ZIP sizes alone are insufficient."""
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > limits["max_xlsx_archive_entries"]:
                raise ParseFailure("XLSX_ENTRY_LIMIT")
            names = [entry.filename for entry in entries]
            if len(set(names)) != len(names):
                raise ParseFailure("XLSX_DUPLICATE_ENTRY")
            total = nodes = text_size = 0
            numbers = {}
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or "\\" in entry.filename
                    or entry.flag_bits & 1
                    or not entry.filename.isascii()
                ):
                    raise ParseFailure("XLSX_UNSAFE_ARCHIVE")
                lowered = entry.filename.lower()
                if "externallinks" in lowered or lowered.endswith((".bin", ".vba")):
                    raise ParseFailure("XLSX_EXTERNAL_CONTENT")
                if entry.file_size > limits["max_xlsx_uncompressed_bytes"]:
                    raise ParseFailure("XLSX_EXPANSION_LIMIT")
                chunks = []
                with archive.open(entry) as stream:
                    while chunk := stream.read(65536):
                        total += len(chunk)
                        if total > limits["max_xlsx_uncompressed_bytes"]:
                            raise ParseFailure("XLSX_EXPANSION_LIMIT")
                        chunks.append(chunk)
                if not lowered.endswith((".xml", ".rels")):
                    continue
                raw_numbers = {}
                for _, element in SafeXML.iterparse(io.BytesIO(b"".join(chunks)), events=("end",)):
                    nodes += 1
                    text_size += len(element.text or "")
                    if nodes > 500000 or text_size > limits["max_parsed_import_bytes"]:
                        raise ParseFailure("XLSX_STRUCTURE_LIMIT")
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag == "Relationship" and element.get("TargetMode") == "External":
                        raise ParseFailure("XLSX_EXTERNAL_CONTENT")
                    if tag == "f":
                        raise ParseFailure("XLSX_FORMULA_UNSUPPORTED")
                    if tag == "c":
                        coordinate = element.get("r", "")
                        match = re.fullmatch(r"([A-Z]{1,3})([0-9]{1,7})", coordinate)
                        if match is None or int(match[2]) > limits["max_import_rows"] + 1:
                            raise ParseFailure("ROW_LIMIT")
                        column = 0
                        for char in match[1]:
                            column = column * 26 + ord(char) - 64
                        if column > limits["max_import_columns"]:
                            raise ParseFailure("COLUMN_LIMIT")
                        value = element.find("{*}v")
                        if element.get("t", "n") == "n" and value is not None and value.text:
                            try:
                                numeric = Decimal(value.text)
                                if not numeric.is_finite() or abs(numeric.adjusted()) > 100:
                                    raise InvalidOperation
                                raw_numbers[coordinate] = format(numeric, "f")
                            except InvalidOperation:
                                raise ParseFailure("XLSX_INVALID_NUMBER") from None
                        element.clear()
                if lowered.startswith("xl/worksheets/"):
                    numbers[entry.filename] = raw_numbers
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise ParseFailure("INVALID_XLSX")
            return numbers
    except ParseFailure:
        raise
    except Exception:
        raise ParseFailure("INVALID_XLSX") from None


def xlsx_rows(content, sheet, limits):
    numbers = inspect_xlsx(content, limits)
    workbook = None
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        sheets = workbook.sheetnames
        if not sheet:
            if len(sheets) != 1:
                raise ParseFailure("SHEET_SELECTION_REQUIRED")
            sheet = sheets[0]
        if sheet not in sheets:
            raise ParseFailure("SHEET_NOT_FOUND")
        worksheet = workbook[sheet]
        # Dimensions are supplied by the upload; never trust them as iteration bounds.
        worksheet.reset_dimensions()
        iterator = worksheet.iter_rows()
        first = next(iterator, ())
        columns = [cell.value for cell in first]
        check_columns(columns, limits)
        numeric = numbers.get(worksheet._worksheet_path.lstrip("/"), {})
        rows = []
        for cells in iterator:
            if len(rows) >= limits["max_import_rows"]:
                raise ParseFailure("ROW_LIMIT")
            if len(cells) > len(columns) and any(
                c.value is not None for c in cells[len(columns) :]
            ):
                raise ParseFailure("XLSX_ROW_WIDTH_INVALID")
            values = []
            for cell in cells[: len(columns)]:
                value = cell.value
                if cell.data_type in {"f", "e"}:
                    raise ParseFailure("XLSX_UNSUPPORTED_CELL")
                check_cell(value, limits)
                if isinstance(value, datetime):
                    if value.time().isoformat() != "00:00:00":
                        raise ParseFailure("DATE_TIME_UNSUPPORTED")
                    value = value.date().isoformat()
                elif isinstance(value, date):
                    value = value.isoformat()
                elif isinstance(value, (int, float)):
                    value = numeric.get(cell.coordinate)
                    if value is None:
                        raise ParseFailure("XLSX_INVALID_NUMBER")
                values.append(value)
            values.extend([None] * (len(columns) - len(values)))
            rows.append(dict(zip(columns, values, strict=True)))
        return columns, rows, sheet
    except ParseFailure:
        raise
    except Exception:
        raise ParseFailure("INVALID_XLSX") from None
    finally:
        if workbook is not None:
            workbook.close()


def five_column_rows(content, descriptor, limits):
    columns, raw, sheet = csv_rows(content, limits)
    aliases = {
        "gstin": "supplier_gstin",
        "invoice number": "invoice_number",
        "date": "invoice_date",
        "taxable value": "taxable_value",
        "tax": "total_tax",
        "supplier_gstin": "supplier_gstin",
        "invoice_number": "invoice_number",
        "invoice_date": "invoice_date",
        "taxable_value": "taxable_value",
        "total_tax": "total_tax",
    }
    mapped = {aliases.get(c.strip().lower()): c for c in columns}
    required = {"supplier_gstin", "invoice_number", "invoice_date", "taxable_value", "total_tax"}
    if set(mapped) != required or len(columns) != 5:
        raise ParseFailure("FIVE_COLUMN_HEADERS_INVALID")
    rows = []
    for i, item in enumerate(raw, 1):
        row = {key: item[header] for key, header in mapped.items()}
        try:
            base = Decimal(row["taxable_value"])
            tax = Decimal(row["total_tax"])
            if not base.is_finite() or not tax.is_finite():
                raise ValueError
            gross = format(base + tax, "f")
        except (InvalidOperation, ValueError):
            gross = "INVALID"
        row.update(
            voucher_id="REGISTER-" + str(i),
            recipient_gstin=descriptor["recipient_gstin"],
            document_type="INVOICE",
            other_charges="0.00",
            round_off="0.00",
            gross_total=gross,
            adapter_assumptions=(
                "User selected invoice-only register with no extra charges or "
                "rounding. Tax components were not supplied."
            ),
        )
        rows.append(row)
    return list(rows[0]) if rows else [], rows, sheet


def gst_2b_rows(content, descriptor, limits):
    try:
        data = bounded_json(
            content.decode("utf-8-sig"), limits["max_json_depth"], exact_numbers=True
        )
    except UnicodeError:
        raise ParseFailure("UTF8_REQUIRED") from None
    if not isinstance(data, dict):
        raise ParseFailure("GST_JSON_SCHEMA_INVALID")
    payload = data.get("data", data)
    if not isinstance(payload, dict) or payload.get("gstin") != descriptor["recipient_gstin"]:
        raise ParseFailure("CONTEXT_MISMATCH")
    rtn = payload.get("rtnprd")
    period = descriptor["period"]
    if rtn not in {period, period[5:] + period[:4]}:
        raise ParseFailure("CONTEXT_MISMATCH")
    docs = payload.get("docdata")
    if not isinstance(docs, dict) or not isinstance(docs.get("b2b"), list):
        raise ParseFailure("GST_JSON_SCHEMA_INVALID")
    # Refuse unsupported document categories rather than silently losing them.
    if any(value for key, value in docs.items() if key != "b2b"):
        raise ParseFailure("GST_JSON_CATEGORY_UNSUPPORTED")
    rows = []

    def amount(value):
        if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
            return None
        try:
            number = Decimal(str(value))
            if not number.is_finite():
                return None
            return format(number, "f")
        except InvalidOperation:
            return None

    for supplier in docs["b2b"]:
        if not isinstance(supplier, dict) or not isinstance(supplier.get("inv"), list):
            raise ParseFailure("GST_JSON_SCHEMA_INVALID")
        for invoice in supplier["inv"]:
            if not isinstance(invoice, dict):
                raise ParseFailure("GST_JSON_SCHEMA_INVALID")
            if len(rows) >= limits["max_import_rows"]:
                raise ParseFailure("ROW_LIMIT")
            values = invoice
            items = invoice.get("itms")
            if items is not None:
                if (
                    not isinstance(items, list)
                    or not items
                    or len(items) > limits["max_import_rows"]
                ):
                    raise ParseFailure("GST_JSON_SCHEMA_INVALID")
                details = [x.get("itm_det", x) if isinstance(x, dict) else None for x in items]
                if any(not isinstance(x, dict) for x in details):
                    raise ParseFailure("GST_JSON_SCHEMA_INVALID")
                values = {}
                for field in ["txval", "iamt", "camt", "samt", "csamt"]:
                    components = [amount(x.get(field)) for x in details]
                    values[field] = (
                        format(sum(Decimal(v) for v in components), "f")
                        if all(v is not None for v in components)
                        else None
                    )
            row = {
                "recipient_gstin": payload["gstin"],
                "supplier_gstin": supplier.get("ctin"),
                "supplier_name": supplier.get("trdnm"),
                "invoice_number": invoice.get("inum"),
                "invoice_date": invoice.get("idt", invoice.get("dt")),
                "document_type": "INVOICE",
                "taxable_value": amount(values.get("txval")),
                "gross_total": amount(invoice.get("val")),
                "igst": amount(values.get("igst", values.get("iamt"))),
                "cgst": amount(values.get("cgst", values.get("camt"))),
                "sgst": amount(values.get("sgst", values.get("samt"))),
                "cess": amount(values.get("cess", values.get("csamt"))),
                "other_charges": "0.00",
                "round_off": "0.00",
                "portal_itc_availability": invoice.get("itcavl", "UNKNOWN"),
            }
            if isinstance(row["invoice_date"], str):
                with suppress(ValueError):
                    row["invoice_date"] = (
                        datetime.strptime(row["invoice_date"], "%d-%m-%Y").date().isoformat()
                    )
            for value in row.values():
                check_cell(value, limits)
            rows.append(row)
    return list(rows[0]) if rows else [], rows, None


def parse_import(content: bytes, descriptor: dict, limits: dict) -> dict:
    if not content or len(content) > limits["max_upload_bytes"]:
        raise ParseFailure("UPLOAD_SIZE_INVALID")
    generated_at = None
    if descriptor["adapter_version"] == "canonical-demo-v1":
        try:
            data = bounded_json(content.decode("utf-8-sig"), limits["max_json_depth"])
        except UnicodeError:
            raise ParseFailure("UTF8_REQUIRED") from None
        if not isinstance(data, dict) or set(data) != {
            "format",
            "provenance",
            "recipient_gstin",
            "period",
            "generated_at",
            "documents",
        }:
            raise ParseFailure("DEMO_JSON_SCHEMA_INVALID")
        if (
            data["format"] != "canonical-demo-v1"
            or data["provenance"] != "SYNTHETIC_DEMO"
            or data["recipient_gstin"] != descriptor["recipient_gstin"]
            or data["period"] != descriptor["period"]
        ):
            raise ParseFailure("CONTEXT_MISMATCH")
        generated_at = data["generated_at"]
        try:
            if not isinstance(generated_at, str) or len(generated_at) > 40:
                raise ValueError
            timestamp = datetime.fromisoformat(generated_at.replace("Z", "+00:00"))
            if timestamp.tzinfo is None:
                raise ValueError
        except ValueError:
            raise ParseFailure("INVALID_GENERATED_AT") from None
        rows = data["documents"]
        if not isinstance(rows, list) or len(rows) > limits["max_import_rows"]:
            raise ParseFailure("ROW_LIMIT")
        columns = []
        for row in rows:
            if not isinstance(row, dict) or not set(row) <= FIELDS:
                raise ParseFailure("DEMO_JSON_SCHEMA_INVALID")
            for key, value in row.items():
                check_cell(value, limits)
                if key not in columns:
                    columns.append(key)
        check_columns(columns, limits)
        sheet = None
    elif descriptor["adapter_version"] == "five-column-v1":
        columns, rows, sheet = five_column_rows(content, descriptor, limits)
    elif descriptor["adapter_version"] == "gst-2b-json-v1":
        columns, rows, sheet = gst_2b_rows(content, descriptor, limits)
    elif descriptor["adapter_version"] == "csv-v1":
        columns, rows, sheet = csv_rows(content, limits)
    elif descriptor["adapter_version"] == "xlsx-v1":
        columns, rows, sheet = xlsx_rows(content, descriptor.get("sheet_name"), limits)
    else:
        raise ParseFailure("UNSUPPORTED_ADAPTER")
    if not rows:
        raise ParseFailure("EMPTY_IMPORT")
    mapping = descriptor.get("mapping") or {field: field for field in columns if field in FIELDS}
    if not set(mapping) <= FIELDS or any(header not in columns for header in mapping.values()):
        raise ParseFailure("INVALID_MAPPING")
    if len(set(mapping.values())) != len(mapping):
        raise ParseFailure("AMBIGUOUS_MAPPING")
    required = REQUIRED | (
        {"voucher_id", "recipient_gstin"} if descriptor["kind"] == "PURCHASE" else set()
    )
    errors = [
        {"field": field, "reason": "MAPPING_REQUIRED"}
        for field in sorted(required - mapping.keys())
    ]
    parsed = [
        canonical_row(row, mapping, descriptor["recipient_gstin"], descriptor["kind"], index)
        for index, row in enumerate(rows, 1)
    ]
    for row in parsed:
        if descriptor["adapter_version"] == "gst-2b-json-v1" and row["original"].get(
            "portal_itc_availability"
        ) not in {"Y", "Yes", "YES", "UNKNOWN"}:
            row["errors"].append({"field": "total_tax", "reason": "PORTAL_ITC_AVAILABILITY_REVIEW"})
        for field in ("period",):
            header = mapping.get(field)
            if header and row["original"].get(header) != descriptor["period"]:
                row["errors"].append({"field": field, "reason": "CONTEXT_MISMATCH"})
        row["accepted"] = not row["errors"] and not errors
    duplicates = flag_duplicates(parsed, descriptor["kind"])
    result = {
        "columns": columns,
        "rows": parsed,
        "mapping": mapping,
        "sheet_name": sheet,
        "errors": errors,
        "generated_at": generated_at,
        "duplicate_rows": duplicates,
    }
    if len(json.dumps(result, ensure_ascii=True).encode()) > limits["max_parsed_import_bytes"]:
        raise ParseFailure("PARSED_RESULT_LIMIT")
    return result
